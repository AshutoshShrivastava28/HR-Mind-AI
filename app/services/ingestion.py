from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.services.upload_validation import SUPPORTED, validate_file


def load_file(path: Path) -> list[Document]:
    """Load a validated PDF, TXT, Markdown, or DOCX file."""
    validate_file(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()

    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), encoding="utf-8").load()

    if suffix == ".docx":
        doc = DocxDocument(str(path))

        paragraphs = [
            paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()
        ]

        table_rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in doc.tables
            for row in table.rows
            if any(cell.text.strip() for cell in row.cells)
        ]

        text = "\n".join(paragraphs + table_rows)

        if not text.strip():
            raise ValueError("No readable text was found in this Word document.")

        return [
            Document(
                page_content=text,
                metadata={"source": path.name},
            )
        ]

    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900,
        chunk_overlap=120,
        add_start_index=True,
    )

    return splitter.split_documents(list(docs))
